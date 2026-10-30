"""PI Reception Activity: the practice calls (Day 3's Reception Activity, replayed on Day 4 before the mock calls).

Five reception calls at a personal injury firm, Harbor & Pine Injury Law (the fictional Florida firm of the Practice
Lab, js/ft-firms.js), each with a challenge a receptionist meets on the line. They are VA Oscar's five reception
recordings redone for the firm (fictional names and numbers): a client who gives the wrong birth year and never got an
appointment email, a frustrated repeat caller, a provider asking for a fax number without saying who it's for, a short
confirmation call, and a client whose ride fell through (a warm transfer). Trainees listen and extract the five data points of each call.

Run:  pip install edge-tts numpy python-docx   then   python3 build/calls/pi_reception_calls.py      (from the repo root)
Writes ft/day3/audio/pi-call-1.mp3 … pi-call-5.mp3, build/calls/pi-reception-calls.json (what the Day 3 and Day 4
deck slides and the trainer's key are built from: build/day-decks.cjs; run it after this) and the trainee's worksheet,
handouts/LSH_Day3_PI_Reception_Activity_Worksheet.docx (needs python-docx). Needs ffmpeg. The voices are Microsoft's
neural text-to-speech (edge-tts); each line is cached in build/calls/.cache (not committed), so a re-run after a
script change only fetches the lines that changed. The noise, line effects and phone tones are made here (numpy).
All names, numbers, dates and claim numbers are fictional; phone numbers use the 555-01xx range.
"""
import asyncio, hashlib, json, os, subprocess, sys
import numpy as np

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
CACHE = os.path.join(REPO, "build", "calls", ".cache")
OUT = os.path.join(REPO, "ft", "day3", "audio")
SR = 16000
RNG = np.random.default_rng(7)

FIRM = "Harbor & Pine Injury Law"
GREET = "Harbor and Pine Injury Law, this is {va}. How may I help you today?"

# ---------- the calls ----------
# a line: (who, text, options). who: "va", "c" (the caller) or another key of the call's voices.
# options: gap (seconds of silence before the line; negative overlaps the line before it), rate / pitch (edge-tts),
# fx (extra effects on this line: "drop" packet loss, "far" caller away from the phone), vol.
# events: ("ring",), ("pickup",), ("hold", secs), ("drop",), ("redial", "5550142290"), ("bed", name) / ("bed", None)
# switch the caller's background on or off; ("pause", secs).
CALLS = [
{
  "id": 1, "title": "The Missing Email", "challenge": "Mixed-up details",
  "tags": ["Mixed-up details: a wrong birth year", "Existing client"],
  "about": "An older client never got the email about the doctor's appointment her case manager set up, and gives the wrong birth year at first.",
  "voices": {"va": ("en-PH-JamesNeural", "Oscar"), "c": ("en-US-EmmaNeural", "Gloria Benton")},
  "line": "home",
  "script": [
    ("ring",), ("pickup",), ("bed", "quiet"),
    ("va", GREET),
    ("c", "Hi, yes, my name is Gloria Benton. I have an appointment with a doctor, with, oh, I don't know if I can say it. It's a long name. Doctor Her... Herrera?", {"rate": "-12%"}),
    ("va", "That's all right, Ms. Benton, I'll look that up for you. First I just need to verify your file. Can you give me your date of birth?"),
    ("c", "July fifth, nineteen fifty.", {"rate": "-12%"}),
    ("va", "Thank you. Hmm, I'm not finding a match with nineteen fifty. Could you give me the year once more?"),
    ("c", "Oh! Sixty. Nineteen sixty. I'm sorry, I said fifty. Nineteen six-zero.", {"rate": "-10%"}),
    ("va", "No problem at all. July fifth, nineteen sixty. And the date of your accident?"),
    ("c", "September third. This year.", {"rate": "-12%"}),
    ("va", "Thank you. Gloria Benton, date of birth July fifth, nineteen sixty, accident September third, twenty twenty-six. You're verified. Now, one moment while I pull up your appointment."),
    ("pause", 1.8),
    ("va", "Okay. I see your appointment with Doctor Herrera at Bayview Orthopedics on Thursday, November twelfth, at one PM. Your case manager, Mara, set that up."),
    ("c", "Yes. That's the one. She said they would send me an email with all the particulars, the address and all that, but I never got it.", {"rate": "-10%"}),
    ("va", "I'm sorry about that. Let me check the email we have for you. I have g benton six zero at mail dot com. Is that right?"),
    ("c", "No, no. It's g dot benton, with a dot. G, dot, B, E, N, T, O, N, six zero, at mail dot com.", {"rate": "-14%"}),
    ("va", "Thank you. So that's g, dot, benton, six zero, at mail dot com. That explains the missing email. I'll note the correct address on your file. In the meantime, would you like the clinic's address now?"),
    ("c", "Yes, please. Let me get a pen. Okay.", {"rate": "-10%"}),
    ("va", "Bayview Orthopedics, twenty-one fifty Bayshore Boulevard, suite three hundred, in Tampa. Their phone is five five five, zero one six, seven seven zero zero.", {"rate": "-10%"}),
    ("c", "Twenty-one fifty Bayshore, suite three hundred. Okay, I have it.", {"rate": "-12%"}),
    ("va", "Perfect. I'll also ask Mara to resend the confirmation to your correct email today. Is this the best number to reach you?"),
    ("c", "Yes, this is my home phone. Five five five, oh one nine, two oh four seven.", {"rate": "-12%"}),
    ("va", "Five five five, zero one nine, two zero four seven. Is there anything else I can help you with, Ms. Benton?"),
    ("c", "No, that's all, dear. Thank you so much.", {"rate": "-10%"}),
    ("va", "You're welcome, Ms. Benton. Take care. Goodbye."),
    ("bed", None), ("hangup",)
  ],
  "key": [
    ("Caller", "Gloria Benton, existing client"),
    ("Callback number", "(555) 019-2047 (home)"),
    ("Verified", "Yes: full name, DOB 07/05/1960 (she said 1950 first; the VA asked again), DOL 09/03/2026"),
    ("Reason for the call", "Never got the confirmation email for her appointment with Dr. Herrera, Bayview Orthopedics, Thu 11/12/2026 at 1:00 PM; the email on file was missing a dot"),
    ("Action taken", "Gave her the clinic's address and phone; corrected email (g.benton60@mail.com) noted on the file; message to Mara Quintero (Case Manager) to resend the confirmation today"),
  ],
  "handled": "Didn't accept a birth year that didn't match: asked again instead of guessing, then found why the email never arrived and gave her the address on the call."
},
{
  "id": 2, "title": "Third Time Calling", "challenge": "A frustrated repeat caller",
  "tags": ["Frustrated repeat caller", "Existing client"],
  "about": "A client in pain calls again about the letter his pain clinic needs before it will book his injection.",
  "voices": {"va": ("en-PH-JamesNeural", "Oscar"), "c": ("en-US-EricNeural", "Andre Wallace")},
  "line": "home",
  "script": [
    ("ring",), ("pickup",), ("bed", "quiet"),
    ("va", GREET),
    ("c", "Yeah, I'm calling to get an update on the letter for my pain doctor. I called last night. I called yesterday morning too. This is the third time.", {"rate": "+10%"}),
    ("va", "I'm sorry you've had to call more than once. Let me check on that for you right now. Can I have your full name and date of birth?"),
    ("c", "Andre Wallace. Two fifteen eighty-five.", {"rate": "+12%"}),
    ("va", "Thank you, Mr. Wallace. That's February fifteenth, nineteen eighty-five?"),
    ("c", "Yes.", {"rate": "+5%"}),
    ("va", "And the date of the accident?"),
    ("c", "July twenty-second.", {"rate": "+8%"}),
    ("va", "July twenty-second, twenty twenty-six. You're verified. One moment while I look at your file."),
    ("pause", 2.0),
    ("va", "Okay. I see the request for a letter of protection to Gulf Coast Pain Center. It's been assigned to our records coordinator, Devin Shaw, and it's due tomorrow."),
    ("c", "Tomorrow? They won't book my injection until they get that letter. I'm in pain every day. I can't keep waiting on this.", {"rate": "+12%"}),
    ("va", "I understand, Mr. Wallace, and I'm sorry you're in pain. Would you like me to try to connect you with Devin now?"),
    ("c", "Yes. Please.", {"rate": "+5%"}),
    ("va", "Please hold for one moment."),
    ("hold", 4.0),
    ("va", "Thank you for holding, Mr. Wallace. Devin is on another call right now. I've sent him an urgent message: that this is your third call, that the clinic won't book your injection without the letter, and that you need it sooner than tomorrow. I've asked him to call you back today. What's the best number for him?"),
    ("c", "This one. Five five five, oh one two, five five eight three.", {"rate": "+8%"}),
    ("va", "Five five five, zero one two, five five eight three. And is there a time that's better for you?"),
    ("c", "Anytime. I'm home.", {"rate": "+5%"}),
    ("va", "Got it. I've noted all of this on your file, so whoever you speak with next will know the whole story. Is there anything else I can help you with?"),
    ("c", "No. Just, please, get it done.", {"rate": "+5%"}),
    ("va", "I understand. Devin will call you today. Thank you for your patience, Mr. Wallace. Goodbye."),
    ("bed", None), ("hangup",)
  ],
  "key": [
    ("Caller", "Andre Wallace, existing client"),
    ("Callback number", "(555) 012-5583, any time"),
    ("Verified", "Yes: full name, DOB 02/15/1985 (read back, as he said it fast), DOL 07/22/2026"),
    ("Reason for the call", "His third call about the Letter of Protection for Gulf Coast Pain Center, which won't book his injection without it; the request is with Devin Shaw, due tomorrow"),
    ("Action taken", "Tried a warm transfer to Devin Shaw (Records Coordinator, on another call); urgent message for a callback today; the repeat calls noted on the file"),
  ],
  "handled": "Apologized once and acted: checked the file at once, tried to reach the person handling it, and passed on an urgent message with the whole story."
},
{
  "id": 3, "title": "Which Office?", "challenge": "An unclear request",
  "tags": ["Unclear request", "Medical provider"],
  "about": "An imaging centre's front desk wants a fax number for the firm's other office and doesn't say who the records are for.",
  "voices": {"va": ("en-PH-JamesNeural", "Oscar"), "c": ("en-US-AvaNeural", "Kayla Jensen")},
  "line": "office",
  "script": [
    ("ring",), ("pickup",), ("bed", "office"),
    ("va", GREET),
    ("c", "Hi, I need the phone number for your Brandon office. I'm faxing something over and I want to make sure they get it."),
    ("va", "I can help with that. All of our offices' calls come here, and the faxes for every office go to our records department. What fax number are you sending it to?"),
    ("c", "It looks like I have five five five, oh one oh, four four one oh."),
    ("va", "Yes, five five five, zero one zero, four four one zero. That's the correct fax number."),
    ("c", "Okay, great. Can I put it to the attention of someone at the Brandon office?"),
    ("va", "Sure. May I ask who's calling, and what the fax is about?"),
    ("c", "Oh, sorry, I didn't say. This is Kayla, from the front desk at Sun City Imaging. It's an MRI report for one of our patients. Tomas Reyes."),
    ("va", "Thank you, Kayla. I'm not able to confirm anything about a file on this call, but you can send it to the attention of Devin Shaw, our records coordinator. Please put the patient's full name and date of birth on the cover sheet, so it gets to the right place."),
    ("c", "Devin Shaw. S, H, A, W?"),
    ("va", "That's right. And may I have your last name and a callback number, in case Devin has a question about the report?"),
    ("c", "Jensen. J, E, N, S, E, N. And our number is five five five, oh one one, six six three zero."),
    ("va", "Kayla Jensen, Sun City Imaging, five five five, zero one one, six six three zero. I'll let Devin know the report is on its way. Is there anything else I can help you with?"),
    ("c", "No, that's it. Thank you!"),
    ("va", "Thank you for calling, Kayla. Have a good day."),
    ("bed", None), ("hangup",)
  ],
  "key": [
    ("Caller", "Kayla Jensen, front desk, Sun City Imaging (a medical provider)"),
    ("Callback number", "(555) 011-6630; the firm's fax confirmed: (555) 010-4410"),
    ("Verified", "A provider, not the client: the patient named as Tomas Reyes; nothing about a file confirmed"),
    ("Reason for the call", "Faxing an MRI report; asked for the Brandon office's number (all offices' calls come to reception and all faxes go to the records department)"),
    ("Action taken", "Confirmed the fax number; told her to fax it attention Devin Shaw (Records Coordinator) with the patient's name and DOB on the cover sheet; Devin told to watch for it"),
  ],
  "handled": "Asked who was calling and what the fax was about before routing it, confirmed the number she had, and confirmed nothing about the patient."
},
{
  "id": 4, "title": "Just Confirming", "challenge": "A short call",
  "tags": ["Short call: every step still counts", "Existing client"],
  "about": "A client calls to confirm tomorrow's meeting with her attorney: a quick call where it's easy to skip a step.",
  "voices": {"va": ("en-PH-JamesNeural", "Oscar"), "c": ("en-US-MichelleNeural", "Nicole Avery")},
  "line": "home",
  "script": [
    ("ring",), ("pickup",), ("bed", "quiet"),
    ("va", GREET),
    ("c", "Hi, I'm just calling to confirm my appointment for tomorrow."),
    ("va", "Sure, I can check that for you. Can I have your full name and date of birth?"),
    ("c", "Nicole Avery. May twenty-eighth, nineteen eighty-three."),
    ("va", "Thank you. And the date of the accident?"),
    ("c", "April ninth."),
    ("va", "Nicole Avery, May twenty-eighth, nineteen eighty-three, accident April ninth, twenty twenty-six. You're verified. One moment."),
    ("pause", 1.4),
    ("va", "I show your meeting with Attorney Celeste Harbor tomorrow, Thursday, October twenty-ninth, at four PM, here at the office."),
    ("c", "Yep, that's it."),
    ("va", "Perfect. And is five five five, zero one three, nine one two six still the best number for you?"),
    ("c", "Yes, it is."),
    ("va", "Great. I've noted your call on the file. Is there anything else I can help you with?"),
    ("c", "No, that's all. Thank you."),
    ("va", "You're welcome, Ms. Avery. See you tomorrow. Goodbye."),
    ("bed", None), ("hangup",)
  ],
  "key": [
    ("Caller", "Nicole Avery, existing client"),
    ("Callback number", "(555) 013-9126 (confirmed as still current)"),
    ("Verified", "Yes: full name, DOB 05/28/1983, DOL 04/09/2026"),
    ("Reason for the call", "To confirm her meeting with Atty. Celeste Harbor tomorrow, Thu 10/29/2026 at 4:00 PM, at the office"),
    ("Action taken", "Confirmed the meeting from the calendar; the call noted on the file"),
  ],
  "handled": "Kept it quick without cutting corners: verified her before reading out the appointment, and checked the callback number on file."
},
{
  "id": 5, "title": "No Ride Tomorrow", "challenge": "Rescheduling and a warm transfer",
  "tags": ["Schedule change and warm transfer", "Existing client"],
  "about": "A client whose usual ride fell through can't get to tomorrow's doctor's appointment and asks to reschedule.",
  "voices": {"va": ("en-PH-JamesNeural", "Oscar"), "c": ("en-US-BrianNeural", "Darnell Pryce"), "mara": ("en-US-JennyNeural", "Mara Quintero")},
  "line": "home",
  "script": [
    ("ring",), ("pickup",), ("bed", "quiet"),
    ("va", GREET),
    ("c", "Hi, uh, good afternoon. My name is Darnell Pryce. I'm a client. I'm calling about my, um, my doctor's appointment tomorrow.", {"rate": "-6%"}),
    ("va", "Good afternoon, Mr. Pryce. I can help with that. First, can I have your date of birth?"),
    ("c", "March thirteenth, nineteen ninety-four."),
    ("va", "Thank you. And the date of the accident?"),
    ("c", "May thirtieth. This year."),
    ("va", "Darnell Pryce, March thirteenth, nineteen ninety-four, accident May thirtieth, twenty twenty-six. You're verified. What's going on with tomorrow's appointment?"),
    ("c", "So usually the veterans' shuttle takes me, but, uh, they just called and nobody's available to drive me tomorrow. And I don't have a car right now, since the accident. Is there any way I can reschedule?", {"rate": "-4%"}),
    ("va", "I'm sorry about that. I see your follow-up with Doctor Herrera at Bayview Orthopedics tomorrow, October twenty-ninth, at ten thirty AM. Your case manager, Mara Quintero, coordinates your appointments, so let me connect you with her. Before I do, what's the best number in case we get disconnected?"),
    ("c", "Five five five, oh one four, seven seven five one."),
    ("va", "Five five five, zero one four, seven seven five one. Please hold for a moment while I get Mara."),
    ("c", "Okay. Thanks."),
    ("hold", 4.0),
    ("bed", None),
    ("mara", "Hi Oscar, it's Mara.", {"internal": True}),
    ("va", "Hi Mara. I have Darnell Pryce on line one. He's verified: date of birth March thirteenth, nineteen ninety-four, accident May thirtieth, twenty twenty-six. His ride to tomorrow's ten thirty with Doctor Herrera fell through, and he has no car, so he'd like to reschedule. His callback is five five five, zero one four, seven seven five one.", {"internal": True}),
    ("mara", "Thanks. I'll take him, and I'll see if we can set up transportation too.", {"internal": True}),
    ("va", "Mr. Pryce? Thank you for waiting. I have Mara on the line for you.", {"gap": 0.6}),
    ("bed", "quiet"),
    ("mara", "Hi Darnell, it's Mara. Oscar filled me in. Let's get you a new date, and let's talk about getting you a ride."),
    ("c", "Thank you. I really appreciate it."),
    ("bed", None), ("hangup",)
  ],
  "key": [
    ("Caller", "Darnell Pryce, existing client"),
    ("Callback number", "(555) 014-7751, taken before the transfer"),
    ("Verified", "Yes: full name, DOB 03/13/1994, DOL 05/30/2026"),
    ("Reason for the call", "His veterans' shuttle can't take him to tomorrow's follow-up with Dr. Herrera, Bayview Orthopedics (10/29/2026, 10:30 AM), and he has no car; wants to reschedule"),
    ("Action taken", "Warm transfer to Mara Quintero (Case Manager, ext. 610), briefed on the internal line: the verification, the appointment, the reason and the callback number"),
  ],
  "handled": "Took the callback number before the hold, and briefed Mara fully so the client didn't have to tell his story twice."
},
]

# ---------- text to speech ----------
def tts(text, voice, rate="+0%", pitch="+0Hz"):
    os.makedirs(CACHE, exist_ok=True)
    k = hashlib.sha1(f"{voice}|{rate}|{pitch}|{text}".encode()).hexdigest()[:20]
    mp3 = os.path.join(CACHE, k + ".mp3")
    if not os.path.exists(mp3) or os.path.getsize(mp3) == 0:
        if os.environ.get("SSL_CERT_FILE"):   # behind a proxy with its own CA: edge-tts checks certificates with certifi's bundle
            import certifi; certifi.where = lambda: os.environ["SSL_CERT_FILE"]
        import edge_tts
        for attempt in range(4):
            try:
                asyncio.run(edge_tts.Communicate(text, voice, rate=rate, pitch=pitch).save(mp3 + ".part"))
                os.replace(mp3 + ".part", mp3); break
            except Exception as e:
                if attempt == 3: raise
                print("  retrying a line:", e)
    raw = subprocess.run(["ffmpeg", "-v", "error", "-i", mp3, "-ac", "1", "-ar", str(SR), "-f", "f32le", "-"], check=True, capture_output=True).stdout
    a = np.frombuffer(raw, dtype=np.float32).copy()
    # trim the silence edge-tts puts around a line
    nz = np.nonzero(np.abs(a) > 0.01)[0]
    return a[max(0, nz[0] - 400): nz[-1] + 800] if len(nz) else a

# ---------- signal helpers ----------
def ffx(a, filt):
    """run a mono float signal through an ffmpeg audio filter chain"""
    p = subprocess.run(["ffmpeg", "-v", "error", "-f", "f32le", "-ar", str(SR), "-ac", "1", "-i", "-", "-af", filt, "-f", "f32le", "-"],
                       input=a.astype(np.float32).tobytes(), check=True, capture_output=True)
    out = np.frombuffer(p.stdout, dtype=np.float32).copy()
    return out[:len(a)] if len(out) >= len(a) else np.pad(out, (0, len(a) - len(out)))

def norm(a, peak=0.9):
    m = np.max(np.abs(a)) if len(a) else 0
    return a * (peak / m) if m > 0 else a

def sec(n): return int(round(n * SR))
def tone(freqs, dur, vol=0.2):
    t = np.arange(sec(dur)) / SR
    return vol * sum(np.sin(2 * np.pi * f * t) for f in freqs) / len(freqs)
def fade(a, ms=8):
    n = min(len(a) // 2, int(SR * ms / 1000))
    if n: a[:n] *= np.linspace(0, 1, n); a[-n:] *= np.linspace(1, 0, n)
    return a
def silence(d): return np.zeros(sec(d), dtype=np.float32)

PHONE = "highpass=f=300,lowpass=f=3400,acompressor=threshold=0.2:ratio=3"
HEADSET = "highpass=f=120,lowpass=f=6500"

def line_fx(a, fx):
    if fx in ("drop", "drop2"):
        # a mobile line breaking up: packets lost (gaps), a garbled codec in between
        a = ffx(a, "acrusher=bits=6:mix=0.35:samples=2," + PHONE)
        step = sec(0.09); rate = 0.16 if fx == "drop" else 0.3
        for i in range(0, len(a), step):
            if RNG.random() < rate: a[i:i + int(step * RNG.uniform(0.6, 2.2))] *= RNG.uniform(0, 0.08)
        return a
    if fx == "far":
        # a speakerphone in a big room: an echo, faint and muffled
        return ffx(a, "aecho=0.8:0.7:70|160:0.45|0.3,lowpass=f=2200," + PHONE)
    if fx in ("near", "near2"):
        a = ffx(a, "aecho=0.8:0.4:60:0.15," + PHONE)
        if fx == "near2":
            for i in range(sec(1.0), len(a), sec(0.5)):
                if RNG.random() < 0.35: a[i:i + sec(0.18)] *= 0.05
        return a
    return ffx(a, PHONE)

# ---------- the caller's background ----------
def brown(n):
    w = RNG.standard_normal(n).astype(np.float32)
    b = np.cumsum(w); b -= np.convolve(b, np.ones(SR) / SR, mode="same")
    return norm(b, 1.0)

def babble(lines, n, voices=("en-US-GuyNeural", "en-US-EmmaNeural", "en-US-EricNeural", "en-US-AvaNeural", "en-US-BrianNeural", "en-US-JennyNeural")):
    out = np.zeros(n, dtype=np.float32)
    for i, (v, t) in enumerate(zip(voices * 3, lines)):
        s = tts(t, v)
        pos = int(RNG.uniform(0, sec(3)))
        while pos < n:
            seg = s[: n - pos]; out[pos: pos + len(seg)] += seg * RNG.uniform(0.4, 1.0)
            pos += len(s) + sec(RNG.uniform(0.5, 3.0))
    return norm(ffx(out, "lowpass=f=2800,aecho=0.6:0.5:40|90:0.3|0.2"), 1.0)

def bed(kind, n):
    if kind == "street_far":   # the caller has stepped away from the road: the same street, further off
        return 0.4 * bed("street", n)
    if kind == "street":
        b = 0.5 * norm(ffx(brown(n), "lowpass=f=500"), 1.0)
        for _ in range(int(n / SR / 7)):   # trucks passing
            at, d = sec(RNG.uniform(0, n / SR)), sec(RNG.uniform(2.5, 4.5))
            env = np.sin(np.linspace(0, np.pi, d)) ** 2
            seg = env * norm(ffx(RNG.standard_normal(d).astype(np.float32), "lowpass=f=900"), 1.0)
            b[at:at + d] += 1.1 * seg[: len(b[at:at + d])]
        for _ in range(int(n / SR / 25) + 1):   # a horn
            at = sec(RNG.uniform(3, max(4, n / SR - 3))); h = fade(0.6 * np.sign(tone([415, 520], 0.5, 1)) * 0.5)
            b[at:at + len(h)] += h[: len(b[at:at + len(h)])]
        for k in range(6):   # a truck backing up
            at = sec(20 + k * 1.0); bp = fade(tone([1050], 0.5, 0.35))
            if at + len(bp) < n: b[at:at + len(bp)] += bp
        return 0.33 * b
    if kind == "car":
        b = 0.55 * norm(ffx(brown(n), "lowpass=f=350"), 1.0) + 0.12 * norm(ffx(RNG.standard_normal(n).astype(np.float32), "bandpass=f=1200:width_type=h:w=900"), 1.0)
        tick = fade(tone([1800], 0.02, 0.6), 2)
        for t in np.arange(8, 14, 0.55):   # a turn signal
            at = sec(t)
            if at + len(tick) < n: b[at:at + len(tick)] += tick
        return 0.28 * b
    if kind == "quiet":
        return 0.04 * norm(ffx(RNG.standard_normal(n).astype(np.float32), "lowpass=f=1200"), 1.0)
    if kind == "office":
        hum = tone([120, 240], n / SR, 0.3)
        return 0.05 * norm(ffx(RNG.standard_normal(n).astype(np.float32), "lowpass=f=1500"), 1.0) + 0.02 * hum
    if kind == "callcenter":
        talk = babble(["Thank you for calling Coastline Mutual, how can I help you?", "Can I get the policy number please?",
                       "I'm going to transfer you to the total loss department.", "Okay, and what's the date of loss on that claim?",
                       "Let me pull that up for you, one moment.", "Yes, the adjuster will reach out within forty-eight hours."], n)
        keys = np.zeros(n, dtype=np.float32)
        click = fade(RNG.standard_normal(sec(0.012)).astype(np.float32) * 0.5, 1)
        t = 0.0
        while t < n / SR - 0.1:
            t += RNG.exponential(0.18); at = sec(t)
            if at + len(click) < n: keys[at:at + len(click)] += click * RNG.uniform(0.3, 1)
        return 0.32 * talk + 0.12 * norm(ffx(keys, "highpass=f=1500"), 1.0)
    if kind == "house":
        tv = babble(["Good evening, and welcome to the six o'clock news.", "Traffic on the Howard Frankland bridge is backed up this evening.",
                     "Coming up after the break, your seven day forecast.", "Scattered storms are expected across the Bay area tomorrow."],
                    n, ("en-US-RogerNeural", "en-US-EmmaNeural", "en-US-SteffanNeural", "en-US-AriaNeural"))
        b = 0.3 * ffx(tv, "lowpass=f=2400")
        bark = norm(ffx(RNG.standard_normal(sec(0.16)).astype(np.float32), "bandpass=f=700:width_type=h:w=500"), 1.0) * np.exp(-np.linspace(0, 5, sec(0.16)))
        for t in (36.0, 37.0, 39.0, 61.0, 61.4):   # a dog
            at = sec(t)
            if at + len(bark) < n: b[at:at + len(bark)] += 0.5 * bark
        return b
    return np.zeros(n, dtype=np.float32)

# ---------- phone sounds ----------
def ring():   # what the VA hears as the line rings: one US ring, then the pick-up
    r = tone([440, 480], 1.6, 0.12); return np.concatenate([fade(r), silence(0.5)])
def click(): return fade(np.concatenate([RNG.standard_normal(sec(0.02)).astype(np.float32) * 0.3, silence(0.2)]), 1)
def dtmf(digits):
    rows = {"1": 697, "2": 697, "3": 697, "4": 770, "5": 770, "6": 770, "7": 852, "8": 852, "9": 852, "0": 941}
    cols = {"1": 1209, "2": 1336, "3": 1477, "4": 1209, "5": 1336, "6": 1477, "7": 1209, "8": 1336, "9": 1477, "0": 1336}
    return np.concatenate([np.concatenate([fade(tone([rows[d], cols[d]], 0.09, 0.15)), silence(0.07)]) for d in digits if d in rows])
def ringback(k=2): return np.concatenate([np.concatenate([fade(tone([440, 480], 2.0, 0.08)), silence(1.2)]) for _ in range(k)])
def hold_music(d):   # the firm's hold tone: a soft arpeggio
    notes, out, t = [262, 330, 392, 523, 392, 330], [], 0
    while t < d:
        f = notes[len(out) % len(notes)]; n = fade(tone([f, f * 2], 0.35, 0.08) * np.exp(-np.linspace(0, 3, sec(0.35))), 5); out.append(n); t += 0.35
    return ffx(np.concatenate(out), PHONE)

# ---------- one call ----------
def render(call):
    voices = call["voices"]
    va_voice = voices["va"][0]
    va_name = voices["va"][1]
    parts, bed_spans, at, cur_bed, bed_from = [], [], 0, None, 0
    def put(a, gap=0.0, caller_side=False):
        nonlocal at
        start = max(0, at + sec(gap)); parts.append((start, a)); at = max(at, start + len(a))
    for item in call["script"]:
        kind = item[0]
        if kind == "ring": put(ring())
        elif kind == "pickup": put(click())
        elif kind == "hangup": put(click(), 0.3); put(silence(0.6))
        elif kind == "pause": put(silence(item[1]))
        elif kind == "hold": put(hold_music(item[1]), 0.3)
        elif kind == "drop":
            put(click(), 0.2)
        elif kind == "redial": put(dtmf(item[1]), 0.6); put(ringback(2), 0.3); put(click())
        elif kind == "bed":
            if cur_bed: bed_spans.append((cur_bed, bed_from, at))
            cur_bed, bed_from = item[1], at
        else:
            who, text = kind, item[1]; o = item[2] if len(item) > 2 else {}
            voice = voices[who][0]
            t = GREET.format(va=va_name) if text == GREET else text
            a = tts(t, voice, o.get("rate", "+0%"), o.get("pitch", "+0Hz"))
            if who == "va": a = ffx(a, HEADSET) * 0.85
            elif o.get("internal"): a = ffx(a, "highpass=f=200,lowpass=f=4500") * 0.8
            else: a = line_fx(a, o.get("fx"))
            # (a caller away from the phone is faint: the level is set after the line is normalized)
            a = norm(a, 0.8) * o.get("vol", 0.4 if o.get("fx") == "far" and who == "c" else 1.0)
            put(a, o.get("gap", 0.35 if who != "va" else 0.45))
    if cur_bed: bed_spans.append((cur_bed, bed_from, at))
    total = at + sec(0.3)
    mix = np.zeros(total, dtype=np.float32)
    for start, a in parts: mix[start:start + len(a)] += a
    for kind, s, e in bed_spans:
        if e > s:
            b = ffx(bed(kind, e - s), PHONE if kind not in ("quiet",) else "lowpass=f=3400")
            n = min(len(b), sec(0.15)); b[:n] *= np.linspace(0, 1, n); b[-n:] *= np.linspace(1, 0, n)
            mix[s:e] += b
    return norm(mix, 0.89)

# the calls play a little faster than the voices speak (pitch kept): 15% quicker, so a trainee isn't waiting on the line
TEMPO = 1.15

def save(a, path):
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "f32le", "-ar", str(SR), "-ac", "1", "-i", "-",
                    "-af", f"atempo={TEMPO},loudnorm=I=-18:TP=-2:LRA=11", "-ar", "22050", "-c:a", "libmp3lame", "-b:a", "56k", path],
                   input=a.astype(np.float32).tobytes(), check=True)

def transcript(call):
    voices, out = call["voices"], []
    for item in call["script"]:
        if item[0] in voices:
            o = item[2] if len(item) > 2 else {}
            who = voices[item[0]][1]
            if item[0] == "va": who = f"{who} (VA)"
            text = (GREET.format(va=voices["va"][1]).replace("Harbor and Pine", "Harbor & Pine") if item[1] == GREET else item[1])
            out.append({"who": who + (" · internal line" if o.get("internal") else ""), "text": text})
        elif item[0] == "hold": out.append({"who": "", "text": "[on hold]"})
        elif item[0] == "drop": out.append({"who": "", "text": "[the call drops]"})
        elif item[0] == "redial": out.append({"who": "", "text": "[the VA calls back]"})
    return out

def worksheet(meta):
    """the trainee's worksheet (handouts/): a table per call with the 5 data points and the challenge's reflection
    (the obstacle, the VA's response) to fill in"""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    d.add_heading("Call Intake & Data Extraction: 5 Practice Calls", 0)
    d.add_paragraph(f"{FIRM} (fictional) · Day 3, Reception Training")
    d.add_paragraph("Listen to each call on the Day 3 slide and fill in its 5 data points, then reflect on the call's challenge: the obstacle (what made the call difficult) and the VA's response (how the receptionist handled it).")
    d.add_paragraph("Name: ____________________________        Batch: ______________        Date: ______________")
    rows = ["1. Caller (name, and who they are)", "2. Callback number", "3. Verified? (how, or why nothing could be shared)",
            "4. Reason for the call", "5. Action taken (who it went to, and how)",
            "Reflection: the obstacle (what made the call difficult)", "Reflection: the VA's response (how the receptionist handled it)"]
    def shade(cell, hexcolor):
        tc = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tc.append(sh)
    for c in meta:
        d.add_heading(f"Call {c['id']}: {c['title']}  ({c['challenge']})", 2)
        t = d.add_table(rows=len(rows), cols=2); t.style = "Table Grid"
        for i, r in enumerate(rows):
            a, b = t.rows[i].cells
            a.text = r; a.paragraphs[0].runs[0].bold = True; shade(a, "E8EBF5")
            a.width = Pt(170); b.width = Pt(300)
            b.text = "\n"
    d.add_paragraph()
    d.add_paragraph("Naming convention: PI_Reception_Activity (VA's first name). Upload it to your Assessment and Activities Folder.")
    path = os.path.join(REPO, "handouts", "LSH_Day3_PI_Reception_Activity_Worksheet.docx")
    d.save(path); print("wrote", os.path.relpath(path, REPO))

def main():
    os.makedirs(OUT, exist_ok=True)
    only = {int(x) for x in sys.argv[1:]}
    meta = []
    for c in CALLS:
        path = os.path.join(OUT, f"pi-call-{c['id']}.mp3")
        if not only or c["id"] in only:
            print(f"Call {c['id']}: {c['title']}")
            save(render(c), path)
        dur = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", path], capture_output=True, text=True, check=True).stdout)
        meta.append({"id": c["id"], "title": c["title"], "challenge": c["challenge"], "tags": c["tags"], "about": c["about"],
                     "va": c["voices"]["va"][1], "file": f"/ft/day3/audio/pi-call-{c['id']}.mp3", "secs": round(dur),
                     "key": [{"k": k, "v": v} for k, v in c["key"]], "handled": c["handled"], "transcript": transcript(c)})
    json.dump({"firm": FIRM + " (fictional)", "calls": meta}, open(os.path.join(REPO, "build", "calls", "pi-reception-calls.json"), "w", encoding="utf8"), ensure_ascii=False, indent=1)
    worksheet(meta)
    print("build/calls/pi-reception-calls.json:", ", ".join(f"{m['id']}. {m['title']} {m['secs'] // 60}:{m['secs'] % 60:02d}" for m in meta))

if __name__ == "__main__":
    main()
